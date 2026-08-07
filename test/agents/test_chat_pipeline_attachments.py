"""`chat_pipeline._build_message_with_attachments` 单元测试。

对应文件上传功能：本轮携带 `file_ids` 时，把附件从 uploads/ 复制进
workspace/ 根目录（沙箱工具 read_file/run_python 的 cwd 所在），并把文件名/
大小/工具可用相对路径拼进用户发言末尾，供模型感知本轮上传了哪些文件、
应该用什么路径访问。.docx/.xlsx/.pdf 额外验证会离线转换出可读的 .md 副本
（受 `ATTACHMENT_AUTO_CONVERT_DOCUMENTS` 开关控制，默认开启）。
"""
from __future__ import annotations

from unittest.mock import AsyncMock, patch

import pytest

from src.agent_core.agents.chat_pipeline import _build_message_with_attachments
from src.agent_core.workspace.thread_workspace_manager import ThreadWorkspaceManager


def _settings_stub(auto_convert: bool = True):
    stub = type("SettingsStub", (), {"ATTACHMENT_AUTO_CONVERT_DOCUMENTS": auto_convert})()
    return stub


@pytest.fixture
def workspace_manager(tmp_path):
    return ThreadWorkspaceManager(str(tmp_path))


def _seed_upload(workspace_manager, conversation_id, user_id, stored_name, content: bytes = b"hello") -> None:
    workspace = workspace_manager.get_or_create(conversation_id, user_id)
    workspace.write_upload(stored_name, content)


@pytest.mark.asyncio
async def test_no_file_ids_returns_original_message() -> None:
    result = await _build_message_with_attachments("conv-1", "user-1", "你好", None)
    assert result == "你好"


@pytest.mark.asyncio
async def test_empty_file_ids_returns_original_message() -> None:
    result = await _build_message_with_attachments("conv-1", "user-1", "你好", [])
    assert result == "你好"


@pytest.mark.asyncio
async def test_plain_text_file_copied_to_workspace_with_relative_path(workspace_manager) -> None:
    _seed_upload(workspace_manager, "conv-1", "user-1", "file-1_report.txt", b"plain text content")

    file_store = AsyncMock()
    file_store.get_files_by_ids = AsyncMock(return_value=[{
        "id": "file-1", "original_name": "report.txt", "content_type": "text/plain",
        "size_bytes": len(b"plain text content"), "stored_name": "file-1_report.txt",
    }])

    with patch("src.agent_core.agents.chat_pipeline.get_file_store", return_value=file_store), \
         patch("src.agent_core.agents.chat_pipeline.get_thread_workspace_manager", return_value=workspace_manager):
        result = await _build_message_with_attachments("conv-1", "user-1", "帮我看看这个文件", ["file-1"])

    assert result.startswith("帮我看看这个文件")
    assert "report.txt" in result
    assert '相对路径 "report.txt"' in result
    # HTTP 下载路径不应该出现在给模型的提示里——那是前端路由，不是沙箱可读路径
    assert "/conversations/" not in result

    workspace = workspace_manager.get_or_create("conv-1", "user-1")
    assert (workspace.workspace_dir / "report.txt").read_bytes() == b"plain text content"
    file_store.get_files_by_ids.assert_awaited_once_with(["file-1"], "conv-1")


@pytest.mark.asyncio
async def test_unknown_file_ids_returns_original_message() -> None:
    file_store = AsyncMock()
    file_store.get_files_by_ids = AsyncMock(return_value=[])

    with patch("src.agent_core.agents.chat_pipeline.get_file_store", return_value=file_store):
        result = await _build_message_with_attachments("conv-1", "user-1", "你好", ["missing-id"])

    assert result == "你好"


@pytest.mark.asyncio
async def test_docx_attachment_gets_markdown_sidecar(workspace_manager) -> None:
    from docx import Document

    workspace = workspace_manager.get_or_create("conv-1", "user-1")
    doc = Document()
    doc.add_paragraph("会议纪要：Q3 目标已达成")
    doc.save(str(workspace.uploads_dir / "file-2_notes.docx"))

    file_store = AsyncMock()
    file_store.get_files_by_ids = AsyncMock(return_value=[{
        "id": "file-2", "original_name": "notes.docx",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": (workspace.uploads_dir / "file-2_notes.docx").stat().st_size,
        "stored_name": "file-2_notes.docx",
    }])

    with patch("src.agent_core.agents.chat_pipeline.get_file_store", return_value=file_store), \
         patch("src.agent_core.agents.chat_pipeline.get_thread_workspace_manager", return_value=workspace_manager):
        result = await _build_message_with_attachments("conv-1", "user-1", "总结一下", ["file-2"])

    assert '相对路径 "notes.md"' in result
    assert "无法被 read_file 直接读取" in result
    assert (workspace.workspace_dir / "notes.md").read_text(encoding="utf-8") == "会议纪要：Q3 目标已达成"


@pytest.mark.asyncio
async def test_auto_convert_disabled_skips_markdown_conversion(workspace_manager) -> None:
    from docx import Document

    workspace = workspace_manager.get_or_create("conv-1", "user-1")
    doc = Document()
    doc.add_paragraph("会议纪要：Q3 目标已达成")
    doc.save(str(workspace.uploads_dir / "file-4_notes.docx"))

    file_store = AsyncMock()
    file_store.get_files_by_ids = AsyncMock(return_value=[{
        "id": "file-4", "original_name": "notes.docx",
        "content_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "size_bytes": (workspace.uploads_dir / "file-4_notes.docx").stat().st_size,
        "stored_name": "file-4_notes.docx",
    }])

    with patch("src.agent_core.agents.chat_pipeline.get_file_store", return_value=file_store), \
         patch("src.agent_core.agents.chat_pipeline.get_thread_workspace_manager", return_value=workspace_manager), \
         patch("src.agent_core.agents.chat_pipeline.get_settings", return_value=_settings_stub(auto_convert=False)):
        result = await _build_message_with_attachments("conv-1", "user-1", "总结一下", ["file-4"])

    assert "自动转换已关闭" in result
    assert not (workspace.workspace_dir / "notes.md").exists()


@pytest.mark.asyncio
async def test_legacy_doc_format_reports_unsupported_reason(workspace_manager) -> None:
    _seed_upload(workspace_manager, "conv-1", "user-1", "file-3_old.doc", b"\xd0\xcf\x11\xe0fake-ole-bytes")

    file_store = AsyncMock()
    file_store.get_files_by_ids = AsyncMock(return_value=[{
        "id": "file-3", "original_name": "old.doc", "content_type": "application/msword",
        "size_bytes": 20, "stored_name": "file-3_old.doc",
    }])

    with patch("src.agent_core.agents.chat_pipeline.get_file_store", return_value=file_store), \
         patch("src.agent_core.agents.chat_pipeline.get_thread_workspace_manager", return_value=workspace_manager):
        result = await _build_message_with_attachments("conv-1", "user-1", "你好", ["file-3"])

    assert "暂不支持解析该格式" in result
    assert "另存为 .docx" in result
